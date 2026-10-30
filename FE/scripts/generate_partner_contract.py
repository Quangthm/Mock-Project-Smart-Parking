"""Generate the browser's partnership draft from the shared contract copy."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
copy = json.loads((ROOT / "src/lib/partnerContract.json").read_text(encoding="utf-8"))
output = ROOT / "public/templates/smartparking-partnership-agreement-draft.docx"

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(8)
for name, size in (("Title", 17), ("Heading 1", 11.5)):
    styles[name].font.name = "Arial"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True

label = document.add_paragraph("SMARTPARKING  •  DỰ THẢO THAM KHẢO")
label.alignment = WD_ALIGN_PARAGRAPH.CENTER
label.runs[0].bold = True
label.runs[0].font.size = Pt(8.5)

title = document.add_paragraph(copy["title"], style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
document.add_paragraph(copy["intro"])

for item in copy["sections"]:
    document.add_paragraph(item["heading"], style="Heading 1")
    for paragraph in item["paragraphs"]:
        document.add_paragraph(paragraph)

document.add_paragraph("THÔNG TIN CẦN HOÀN THIỆN", style="Heading 1")
for line in (
    "Đối tác / doanh nghiệp: _______________________________________________",
    "Người đại diện có thẩm quyền: ________________________________________",
    "Tên và địa chỉ bãi đỗ: _________________________________________________",
    "Phụ lục thương mại và ngày hiệu lực: __________________________________",
):
    document.add_paragraph(line)

document.add_paragraph("XÁC NHẬN CỦA HAI BÊN", style="Heading 1")
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text = "ĐẠI DIỆN ĐỐI TÁC"
table.cell(0, 1).text = "ĐẠI DIỆN SMARTPARKING"
table.cell(1, 0).text = "Họ tên: ____________________\nNgày: ______________________\nChữ ký: ____________________"
table.cell(1, 1).text = "Họ tên: ____________________\nNgày: ______________________\nChữ ký, dấu: _______________"
for row in table.rows:
    for cell in row.cells:
        cell.width = Inches(3.2)
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        cell._tc.get_or_add_tcPr().append(borders)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(5)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("SmartParking • Dự thảo cần được rà soát và ký chính thức")
footer.runs[0].font.size = Pt(8)

output.parent.mkdir(parents=True, exist_ok=True)
document.save(output)
print(output)
